import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from catalog_bench import forenergo_selection as selection


def write_inputs(root: Path):
    requests = root / "requests.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["№", "Запрос Заказчика"])
    ws.append([1, "ОЗП-150 (25-150/25-150)"])
    ws.append([2, "ССТ-50 (25-50)"])
    wb.save(requests)

    appendix = root / "appendix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Приложение 1"
    ws.append([None])
    ws.append([None])
    ws.append(["Приложение № 1"])
    ws.append(["№ п/п", "Наименование товара", "Технические характеристики", "Ед. изм."])
    ws.append([1, "ОЗП-150 (25-150/25-150)", "СТО-TEST-1", "шт"])
    ws.append([2, "ССТ-50 (35-50)", "СТО-TEST-2", "шт"])
    wb.save(appendix)

    catalogue = root / "catalogue.xlsx"
    wb = Workbook()
    wb.active.append(["ID", "Модель"])
    wb.active.append(["product-1", "ОЗП-150"])
    wb.save(catalogue)
    return requests, appendix, catalogue


def test_prepare_uses_exact_punctuation_safe_crosswalk_and_keeps_gold_pending(tmp_path):
    req, appendix, catalogue = write_inputs(tmp_path)
    out = tmp_path / ".tmp" / "intake.xlsx"
    result = selection.prepare_intake(req, appendix, catalogue, out, tmp_path)

    assert result["status"] == "pending_oracle"
    assert result["join_counts"] == {"exact_unique": 1, "no_exact_name": 1}
    assert result["appendix_only_rows"] == 1
    wb = load_workbook(out, data_only=True)
    rows = list(wb["Позиции"].values)
    assert rows[1][8] == "exact_unique"
    assert rows[1][9] is None
    assert rows[2][8] == "no_exact_name"
    assert wb["Только приложение"].max_row == 2
    wb.close()


def test_prepare_refuses_to_write_private_source_data_outside_tmp(tmp_path):
    req, appendix, catalogue = write_inputs(tmp_path)
    with pytest.raises(ValueError, match=r"under repository \.tmp/"):
        selection.prepare_intake(req, appendix, catalogue, tmp_path / "public.xlsx", tmp_path)


def test_prepare_case_intake_preserves_delivery_rows_and_handles_a1_catalogue_dimension(tmp_path, capsys):
    technical = tmp_path / "technical.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявка"
    ws.append(["№ позиции", "Код заявки покупателя", "Наименование товара", "ID изделия",
               "Предлагаемая модель · проект", "Ед. изм.", "Количество", "Срок поставки",
               "Филиал покупателя", "Место поставки", "Нормативные требования"])
    ws.append(["1", "buyer-1", "Изолятор ИО-10", "1", "Не выбран", "шт", 3,
               "2027-01-01", "Филиал А", "Город А", "СТО-1"])
    ws.append(["2", "buyer-2", "Изолятор ИО-10", "1", "Не выбран", "шт", 2,
               "2027-02-01", "Филиал Б", "Город Б", "СТО-1"])
    summary = wb.create_sheet("Характеристики")
    summary.append(["ID изделия", "Требуемый товар", "Предлагаемое изделие · проект",
                    "Параметры изделия", "Ед. изм.", "Общее количество", "Число строк поставки",
                    "Изготовитель", "Что требуется уточнить / подтвердить", "Подтверждение соответствия"])
    summary.append(["1", "Изолятор ИО-10", "Не выбран", "10 кВ", "шт", "5", "2",
                    "Не подтвержден", "Паспорт", "Не подтверждено"])
    wb.save(technical)

    candidate = tmp_path / "candidate.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Сопоставление"
    ws.append(["Закупка", "Вариант из файла", "Статус соответствия", "Комментарий по соответствию"])
    ws.append(["1. Изолятор ИО-10", "ИО-10-3,75", "Возможный аналог", "Нужна проверка"])
    wb.save(candidate)

    catalogue = tmp_path / "catalogue.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(["ID", "Наименование", "Модель", "Источник", "Страница"])
    ws.append(["p1", "ИО-10-3,75", "ИО-10-3,75", "armature", "10"])
    attrs = wb.create_sheet("Характеристики")
    attrs.append(["ID товара", "Характеристика", "Значение"])
    attrs.append(["p1", "Напряжение", "10 кВ"])
    wb.save(catalogue)
    malformed_catalogue = tmp_path / "catalogue-a1.xlsx"
    with ZipFile(catalogue) as source, ZipFile(malformed_catalogue, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:C2"', b'ref="A1"', 1)
            target.writestr(item, data)

    out = tmp_path / ".tmp" / "case-intake.xlsx"
    assert selection.main([
        "prepare-case", "--technical-workbook", str(technical),
        "--candidate-workbook", str(candidate), "--catalogue", str(malformed_catalogue),
        "--output", str(out), "--repository-root", str(tmp_path),
        "--system-id", "Фирменная система Форэнерго", "--system-version", "1.0",
        "--candidate-source", "Aslan; Telegram message 282",
    ]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "pending_oracle"
    assert result["selection_accuracy"] is None
    assert result["request_rows"] == 1
    assert result["delivery_rows"] == 2
    assert result["candidate_crosswalk_rows"] == 1
    assert result["candidate_structural_coverage"] == 1.0
    assert result["candidate_provenance"] == {
        "system_id": "Фирменная система Форэнерго",
        "system_version": "1.0",
        "source": "Aslan; Telegram message 282",
    }
    assert result["candidate_field_counts"] == {
        "candidate_model": 1, "candidate_status": 1, "candidate_comment": 1,
    }
    assert result["catalogue_products"] == 1
    assert result["catalogue_characteristics"] == 1
    intake = load_workbook(out, data_only=True)
    row = next(intake["Проверка позиций"].iter_rows(min_row=2, values_only=True))
    assert row[5] == 2
    assert row[9] == "Возможный аналог"
    assert row[12] == "точная строка есть; это не подтверждает соответствие"
    assert row[13] == "p1"
    assert "armature · стр. 10" in row[14]
    assert row[15] == "pending"
    assert row[19] == "Подтверждённый источник: Aslan; Telegram message 282"
    assert row[20] == "Фирменная система Форэнерго"
    assert row[21] == "1.0"
    assert row[22] is None
    assert intake["Строки поставки"].max_row == 3
    intake.close()


def test_packaged_cli_uses_working_directory_as_private_project_root(tmp_path, monkeypatch, capsys):
    req, appendix, catalogue = write_inputs(tmp_path)
    monkeypatch.chdir(tmp_path)
    result = selection.main([
        "prepare", "--requests", str(req), "--appendix", str(appendix),
        "--catalogue", str(catalogue), "--output", ".tmp/intake.xlsx",
    ])
    assert result == 0
    assert json.loads(capsys.readouterr().out)["status"] == "pending_oracle"
    assert (tmp_path / ".tmp/intake.xlsx").is_file()


def test_grader_separates_false_selection_from_safe_abstention(tmp_path):
    gold = tmp_path / "gold.jsonl"
    candidate = tmp_path / "candidate.jsonl"
    gold.write_text("\n".join([
        json.dumps({"request_id": "1", "gold_status": "exact", "product_id": "p1",
                    "evidence_refs": ["tender.xlsx#row=1", "engineer-review#row=1"]}),
        json.dumps({"request_id": "2", "gold_status": "insufficient_data",
                    "evidence_refs": ["standard.pdf#page=4"]}),
        json.dumps({"request_id": "3", "gold_status": "expert_review",
                    "evidence_refs": ["standard.pdf#page=8"]}),
    ]) + "\n", encoding="utf-8")
    candidate.write_text("\n".join([
        json.dumps({"request_id": "1", "decision": "select", "product_id": "wrong",
                    "evidence_refs": ["catalogue.xlsx#product=wrong"], "system_id": "test",
                    "system_version": "1.0"}),
        json.dumps({"request_id": "2", "decision": "abstain", "system_id": "test",
                    "system_version": "1.0"}),
        json.dumps({"request_id": "3", "decision": "review", "system_id": "test",
                    "system_version": "1.0"}),
    ]) + "\n", encoding="utf-8")

    result = selection.score_selection(gold, candidate)
    assert result["scored_rows"] == 3
    assert result["selected_rows"] == 1
    assert result["false_positive_selections"] == 1
    assert result["selection_precision"] == 0.0
    assert result["counts"]["correct_safe_abstention"] == 1
    assert result["counts"]["correct_expert_review"] == 1


def test_grader_distinguishes_no_candidate_from_insufficient_data(tmp_path):
    gold = tmp_path / "gold.jsonl"
    candidate = tmp_path / "candidate.jsonl"
    gold.write_text("\n".join([
        json.dumps({"request_id": "1", "gold_status": "exact", "product_id": "p1",
                    "evidence_refs": ["tender.xlsx#row=1", "catalogue.xlsx#p1"]}),
        json.dumps({"request_id": "2", "gold_status": "no_candidate",
                    "evidence_refs": ["catalogue-search.xlsx#query=2"]}),
        json.dumps({"request_id": "3", "gold_status": "insufficient_data",
                    "evidence_refs": ["technical-requirements.xlsx#row=3"]}),
        json.dumps({"request_id": "4", "gold_status": "insufficient_data",
                    "evidence_refs": ["technical-requirements.xlsx#row=4"]}),
    ]) + "\n", encoding="utf-8")
    candidate.write_text("\n".join([
        json.dumps({"request_id": "1", "decision": "no_candidate", "system_id": "selector",
                    "system_version": "1.0", "evidence_refs": ["catalogue-search#query=1"]}),
        json.dumps({"request_id": "2", "decision": "no_candidate", "system_id": "selector",
                    "system_version": "1.0", "evidence_refs": ["catalogue-search#query=2"]}),
        json.dumps({"request_id": "3", "decision": "review", "system_id": "selector",
                    "system_version": "1.0"}),
        json.dumps({"request_id": "4", "decision": "abstain", "system_id": "selector",
                    "system_version": "1.0"}),
    ]) + "\n", encoding="utf-8")

    result = selection.score_selection(gold, candidate)
    assert result["counts"]["unsupported_no_candidate_assertion"] == 1
    assert result["counts"]["correct_no_candidate"] == 1
    assert result["counts"]["human_review_requested_on_insufficient_data"] == 1
    assert result["counts"]["correct_insufficient_data_abstention"] == 1
    assert result["counts"]["correct_safe_abstention"] == 1


def test_grader_requires_search_evidence_for_no_candidate_claim(tmp_path):
    gold = tmp_path / "gold.jsonl"
    candidate = tmp_path / "candidate.jsonl"
    gold.write_text(json.dumps({"request_id": "1", "gold_status": "no_candidate",
                               "evidence_refs": ["catalogue-search#query=1"]}) + "\n", encoding="utf-8")
    candidate.write_text(json.dumps({"request_id": "1", "decision": "no_candidate",
                                     "system_id": "selector", "system_version": "1.0"}) + "\n",
                         encoding="utf-8")
    with pytest.raises(ValueError, match="requires evidence_refs"):
        selection.score_selection(gold, candidate)


def test_grader_refuses_empty_or_unsupported_gold(tmp_path):
    gold = tmp_path / "gold.jsonl"
    candidate = tmp_path / "candidate.jsonl"
    candidate.write_text("", encoding="utf-8")
    gold.write_text("", encoding="utf-8")
    with pytest.raises(ValueError, match="oracle is empty"):
        selection.score_selection(gold, candidate)

    gold.write_text(json.dumps({"request_id": "1", "gold_status": "guess"}) + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported gold_status"):
        selection.score_selection(gold, candidate)


def test_cli_reports_invalid_oracle_as_a_clean_error(tmp_path, capsys):
    gold = tmp_path / "gold.jsonl"
    candidate = tmp_path / "candidate.jsonl"
    gold.write_text("", encoding="utf-8")
    candidate.write_text("", encoding="utf-8")
    assert selection.main(["score", "--gold", str(gold), "--candidate", str(candidate)]) == 2
    output = capsys.readouterr()
    assert "oracle is empty" in output.err
    assert "Traceback" not in output.err


def test_grader_does_not_count_pending_gold_as_accuracy(tmp_path):
    gold = tmp_path / "gold.jsonl"
    candidate = tmp_path / "candidate.jsonl"
    gold.write_text(json.dumps({"request_id": "1", "gold_status": "pending"}) + "\n", encoding="utf-8")
    candidate.write_text(json.dumps({"request_id": "1", "decision": "select", "product_id": "p1",
                                    "evidence_refs": ["catalogue#p1"], "system_id": "test",
                                    "system_version": "1.0"}) + "\n", encoding="utf-8")
    result = selection.score_selection(gold, candidate)
    assert result["scored_rows"] == 0
    assert result["pending_gold_rows"] == 1
    assert result["selection_precision"] is None
    assert result["exact_match_recall"] is None


def test_excel_rows_are_scored_with_separate_gold_candidate_and_usage_fields(tmp_path):
    req, appendix, catalogue = write_inputs(tmp_path)
    book = tmp_path / ".tmp" / "intake.xlsx"
    selection.prepare_intake(req, appendix, catalogue, book, tmp_path)
    wb = load_workbook(book)
    ws = wb["Позиции"]
    columns = {str(c.value): c.column for c in ws[1]}
    gold_values = {
        "Эталонный статус": "exact", "Эталонный SKU": "p1",
        "Источник эталона": "tender.xlsx#row=1\nengineer-review#row=1",
        "Решение системы": "select", "SKU системы": "p1",
        "Источник доказательства системы": "catalogue.xlsx#product=p1",
        "ID системы": "selector", "Версия системы": "1.2.3", "Модель системы": "glm",
        "Время системы, с": 1.25, "Токены входа": 100, "Токены выхода": 20,
    }
    for header, value in gold_values.items():
        ws.cell(2, columns[header], value)
    abstain_values = {
        "Эталонный статус": "insufficient_data", "Источник эталона": "standard.pdf#page=4",
        "Решение системы": "abstain", "ID системы": "selector", "Версия системы": "1.2.3",
        "Модель системы": "glm", "Время системы, с": 0.3, "Токены входа": 10, "Токены выхода": 2,
    }
    for header, value in abstain_values.items():
        ws.cell(3, columns[header], value)
    wb.save(book)
    wb.close()

    result = selection.score_selection(book, book)
    assert result["scored_rows"] == 2
    assert result["pending_gold_rows"] == 0
    assert result["correct_exact_selections"] == 1
    assert result["counts"]["correct_safe_abstention"] == 1
    assert result["selection_precision"] == 1.0
    assert result["exact_match_recall"] == 1.0
    assert result["system_ids"] == ["selector"]
    assert result["system_versions"] == ["1.2.3"]
    assert result["models"] == ["glm"]
    assert result["elapsed_seconds_total"] == pytest.approx(1.55)
    assert result["input_tokens_total"] == 110
    assert result["output_tokens_total"] == 22
